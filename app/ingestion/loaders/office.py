import logfire


def parse_office(file_path: str):
    """
    Parses Office documents (.docx, .pptx) using python-docx and python-pptx directly.
    Avoids unstructured/PyTorch which causes DLL crashes on Windows.
    """
    with logfire.span("📄 Office Document Parsing", filename=file_path):
        try:
            lower = file_path.lower()

            if lower.endswith(".docx"):
                from docx import Document
                doc = Document(file_path)
                full_text = "\n\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())

            elif lower.endswith(".pptx"):
                from pptx import Presentation
                prs = Presentation(file_path)
                parts = []
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            parts.append(shape.text.strip())
                full_text = "\n\n".join(parts)

            else:
                raise ValueError(f"Unsupported Office format: {file_path}")

            if not full_text.strip():
                logfire.warning(f"⚠️ Empty text extracted from {file_path}")
            else:
                logfire.info(f"✅ Successfully parsed {len(full_text)} characters")

            return full_text

        except Exception as e:
            logfire.error(f"❌ Office Parse Failed: {e}")
            raise e

